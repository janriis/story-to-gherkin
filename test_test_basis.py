import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from test_basis import TestBasisError, read_test_basis


class ReadTestBasisTests(unittest.TestCase):
    def test_utf8_bom_in_txt_is_not_returned(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "basis.txt"
            path.write_bytes(b"\xef\xbb\xbfKrav fra tekstfil")

            self.assertEqual(read_test_basis(path), "Krav fra tekstfil")

    def test_md_text_is_read_as_utf8_with_bom(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "basis.md"
            path.write_bytes(b"\xef\xbb\xbf" + "# Acceptkriterie\n\nBrugeren kan logge ind.".encode("utf-8"))

            self.assertEqual(
                read_test_basis(path), "# Acceptkriterie\n\nBrugeren kan logge ind."
            )

    def test_pdf_reads_selectable_text_from_both_pages_in_order(self):
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "basis.pdf"
            pdf = canvas.Canvas(str(path))
            pdf.drawString(72, 720, "Krav A")
            pdf.showPage()
            pdf.drawString(72, 720, "Krav B")
            pdf.save()

            text = read_test_basis(path)

        self.assertLess(text.index("Krav A"), text.index("Krav B"))

    def test_docx_preserves_paragraph_table_and_paragraph_order(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "basis.docx"
            document = Document()
            document.add_paragraph("Første afsnit")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Tabel venstre"
            table.cell(0, 1).text = "Tabel højre"
            document.add_paragraph("Sidste afsnit")
            document.save(path)

            text = read_test_basis(path)

        self.assertLess(text.index("Første afsnit"), text.index("Tabel venstre"))
        self.assertLess(text.index("Tabel venstre"), text.index("Sidste afsnit"))
        self.assertIn("Tabel venstre | Tabel højre", text)

    def test_empty_or_scanned_pdf_is_rejected(self):
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scannet.pdf"
            pdf = canvas.Canvas(str(path))
            pdf.showPage()
            pdf.save()

            with self.assertRaisesRegex(TestBasisError, "ingen læsbar tekst"):
                read_test_basis(path)

    def test_empty_docx_is_rejected(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "tom.docx"
            Document().save(path)

            with self.assertRaisesRegex(TestBasisError, "ingen læsbar tekst"):
                read_test_basis(path)

    def test_corrupt_docx_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "korrupt.docx"
            path.write_bytes(b"ikke et Word-dokument")

            with self.assertRaisesRegex(TestBasisError, "kan ikke læses"):
                read_test_basis(path)

    def test_docx_with_malformed_document_xml_is_rejected(self):
        """Catches XML parser errors escaping the document-reader contract."""
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            path = Path(directory) / "malformed.docx"
            document = Document()
            document.add_paragraph("Readable before corruption")
            document.save(source)
            with ZipFile(source) as original, ZipFile(path, "w") as malformed:
                for member in original.infolist():
                    content = (
                        b"<w:document>" if member.filename == "word/document.xml"
                        else original.read(member)
                    )
                    malformed.writestr(member, content)

            with self.assertRaisesRegex(TestBasisError, "kan ikke læses"):
                read_test_basis(path)

    def test_docx_with_only_empty_table_cells_is_rejected(self):
        """Catches table separators being mistaken for readable requirements."""
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty-table.docx"
            document = Document()
            document.add_table(rows=1, cols=2)
            document.save(path)

            with self.assertRaisesRegex(TestBasisError, "ingen læsbar tekst"):
                read_test_basis(path)

    def test_docx_nested_table_text_keeps_body_reading_order(self):
        """Catches silently dropping requirements inside a table cell's table."""
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested.docx"
            document = Document()
            document.add_paragraph("Start")
            outer = document.add_table(rows=1, cols=2)
            first_cell = outer.cell(0, 0)
            first_cell.text = "Before"
            nested = first_cell.add_table(rows=1, cols=2)
            nested.cell(0, 0).text = "Nested left"
            nested.cell(0, 1).text = "Nested right"
            first_cell.add_paragraph("After")
            outer.cell(0, 1).text = "Other cell"
            document.add_paragraph("End")
            document.save(path)

            text = read_test_basis(path)

        positions = [text.index(part) for part in (
            "Start", "Before", "Nested left", "Nested right", "After", "Other cell", "End",
        )]
        self.assertEqual(positions, sorted(positions))

    def test_protected_pdf_is_rejected(self):
        from pypdf import PdfReader, PdfWriter
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.pdf"
            path = Path(directory) / "beskyttet.pdf"
            pdf = canvas.Canvas(str(source))
            pdf.drawString(72, 720, "Hemmelig kravtekst")
            pdf.save()
            writer = PdfWriter()
            writer.append(PdfReader(str(source)))
            writer.encrypt("kodeord")
            with path.open("wb") as stream:
                writer.write(stream)

            with self.assertRaisesRegex(TestBasisError, "beskyttet"):
                read_test_basis(path)

    def test_missing_file_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "findes-ikke-testbasis.txt"

            with self.assertRaisesRegex(TestBasisError, "kan ikke læses"):
                read_test_basis(path)

    def test_doc_and_unknown_extensions_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            for name in ("gammel.doc", "basis.rtf"):
                path = Path(directory) / name
                path.write_text("tekst", encoding="utf-8")

                with self.subTest(path=path), self.assertRaisesRegex(
                    TestBasisError, "understøttes ikke"
                ):
                    read_test_basis(path)


if __name__ == "__main__":
    unittest.main()
